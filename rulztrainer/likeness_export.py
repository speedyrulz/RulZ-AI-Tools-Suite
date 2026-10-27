"""A run's likeness probes as a spreadsheet (.xlsx): the Monitor's Export spreadsheet button.

    Updated scores   one row per probe: the smoothed face, body and combined score (what the Monitor and the best
                     checkpoint use) and each probe's own, as live formulas, with the best smoothed combined score
    Raw probe scores the raw face (ArcFace) and body (DINOv2 / proportions) similarities from likeness.jsonl
    Scale            the floors (the step-0 probe) and the references' own levels from the run, the span / curve and
                     the weights from its config: editable, the scores follow

The scores are computed in the sheet from the raw ones, the trainer's normalize(): min(100, 100 * max(0, (raw - floor)
/ (top - floor)) ^ curve) with top = floor + (ceiling - floor) * span, the body's ceiling at least halfway from its floor
to 100. A run from before a scale change is rescored on the span and curve its config file has now (a config saved
before body_scale_span existed gets the current defaults). In a race each configuration gets its own pair of sheets
(and the winner's probes after the race theirs); each has its own column on the Scale sheet.

build_workbook_from() makes the same workbook from probe records and settings without a job folder: the Previous Runs
tab exports a run from its history record (history.probe_raw keeps its raw scores), whose job folder a later run with
the same job name may have overwritten."""
from __future__ import annotations

import json
from pathlib import Path

from .likeness import CEILING_DEFAULT, MIN_ROOM_SHARE, MIN_SPAN

FONT = "Arial"


def read_probes(job_dir: Path) -> list[dict]:
    """likeness.jsonl's records in order; a probe logged twice (the step-0 probe) is kept once, the later copy."""
    out: list[dict] = []
    p = Path(job_dir) / "likeness.jsonl"
    for line in p.read_text(encoding="utf-8").splitlines() if p.is_file() else []:
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        if out and out[-1].get("step") == rec.get("step") and out[-1].get("images") == rec.get("images"):
            out[-1] = rec
        else:
            out.append(rec)
    return out


def _load_cfg(job_dir: Path):
    from .config import TrainConfig

    try:
        cfg, _, _ = TrainConfig.load_report(Path(job_dir) / "config.yaml")
        return cfg
    except Exception:  # noqa: BLE001 - no or unreadable config: the defaults
        return TrainConfig()


def _race_names(job_dir: Path) -> list[str]:
    try:
        return json.loads((Path(job_dir) / "race.json").read_text(encoding="utf-8")).get("names") or []
    except (OSError, ValueError, AttributeError):
        return []


def _groups(probes: list[dict], names) -> list[tuple[str, list[dict]]]:
    """[(label, probes)]: the run as one group, or a race's configurations (and "winner" after the race). A record
    names its configuration in "racer" (history.probe_raw), else its images do."""
    from .history import probe_racer

    if not names:
        return [("", probes)]
    by: dict[str, list[dict]] = {}
    for pr in probes:
        by.setdefault(pr["racer"] if "racer" in pr else probe_racer(pr, names), []).append(pr)
    return [(n or "winner", by[n]) for n in [*names, ""] if n in by]


def _reference_levels(probes: list[dict], cfg) -> dict:
    """{kind: (floor, the references' own level)} in raw points: the run's logged scale with its stretch taken out
    (the record's span, else the stretch the curve came with before the span was logged), or the step-0 probe and
    the default ceiling for a run without normalization."""
    scale = next((p["scale"] for p in probes if p.get("scale")), None)
    out = {}
    for kind in ("face", "body"):
        if scale and kind in scale:
            lo, top = (float(x) for x in scale[kind])
            span = (scale.get("span") or {}).get(kind)
            if span is None:  # logged before the span was: face stretched with its curve (1.35), body linear
                span = 1.35 if float((scale.get("curve") or {}).get(kind, 1.0)) > 1.0 else 1.0
            out[kind] = (lo, lo + (top - lo) / float(span))
        else:
            first = next((p.get(kind) for p in probes if p.get(kind) is not None), None)
            lo = float(first) if first is not None else 0.0
            out[kind] = (lo, max(CEILING_DEFAULT[kind], lo + MIN_SPAN[kind]))
    return out


def build_workbook(job_dir: Path):
    """The workbook for a job folder (openpyxl Workbook). Raises FileNotFoundError without probes."""
    job_dir = Path(job_dir)
    probes = read_probes(job_dir)
    if not probes:
        raise FileNotFoundError(f"No likeness probes in {job_dir} yet (likeness.jsonl)")
    return build_workbook_from(probes, _load_cfg(job_dir), _race_names(job_dir))


def build_workbook_from(probes: list[dict], cfg, race_names=()):
    """The workbook for probe records ({"step", "face", "body"} raw, "scale" as the trainer logs it, a race's "racer"
    or images) and a config (its job_name, scale span / curve, weights and probe_ema go on the Scale sheet)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if not probes:
        raise FileNotFoundError("No likeness probes to export")
    groups = _groups(probes, race_names)

    plain, bold = Font(name=FONT), Font(name=FONT, bold=True)
    blue, green = Font(name=FONT, color="0000FF"), Font(name=FONT, color="008000")
    note = Font(name=FONT, italic=True, color="6B7280")
    head, head_fill = Font(name=FONT, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")
    line = Border(bottom=Side(style="thin", color="D4D4D8"))

    wb = Workbook()
    wb.remove(wb.active)
    sa = wb.create_sheet("Scale")
    sa["A1"] = f"{cfg.job_name}: likeness scale"
    sa["A1"].font = Font(name=FONT, bold=True, size=12)
    labels = ["Face floor (raw, step-0 probe)", "Face references' own level (raw)", "Face scale span", "Face scale curve",
              "Body floor (raw, step-0 probe)", "Body references' own level (raw)", "Body scale span",
              "Body scale curve", "Body minimum share of the room to 100", "Face weight", "Body weight",
              "Probe smoothing (probe_ema)", "Face scale top (raw)", "Body scale top (raw)",
              "Face: your references' own level reads", "Body: your references' own level reads"]
    notes = ["The run's step-0 probe (its logged scale).", "What your references score against each other.",
             "face_scale_span", "face_scale_curve", "The run's step-0 probe (its logged scale).",
             "What your references score against each other.", "body_scale_span", "body_scale_curve",
             "The body ceiling sits at least this share of the way from its floor to 100.", "face_weight",
             "body_weight", "Weight on the previous smoothed value.",
             "Computed: floor + (references' level - floor) x span.",
             "Computed: the ceiling at least 8 points and the minimum share above the floor, then x span.",
             "Computed: where the references' own face level lands on this scale.",
             "Computed: where the references' own body level lands on this scale."]
    for i, (lab, txt) in enumerate(zip(labels, notes)):
        sa.cell(3 + i, 1, lab).font = bold if i >= 12 else plain
        sa.cell(3 + i, 2 + len(groups), txt).font = note
    sa.column_dimensions["A"].width = 40
    sa.column_dimensions[get_column_letter(2 + len(groups))].width = 70

    score_sheets, raw_sheets = [], []
    for g, (label, rows) in enumerate(groups):
        col = get_column_letter(2 + g)
        sa.cell(2, 2 + g, label or "run").font = bold
        sa.column_dimensions[col].width = 13
        levels = _reference_levels(rows, cfg)
        values = [levels["face"][0], levels["face"][1], cfg.face_scale_span, cfg.face_scale_curve, levels["body"][0],
                  levels["body"][1], getattr(cfg, "body_scale_span", 1.35), getattr(cfg, "body_scale_curve", 1.5),
                  MIN_ROOM_SHARE["body"], cfg.face_weight, cfg.body_weight, cfg.probe_ema]
        for i, v in enumerate(values):
            c = sa.cell(3 + i, 2 + g, float(v))
            c.font, c.number_format = blue, "0.00"
        sa.cell(15, 2 + g, f"={col}3+({col}4-{col}3)*{col}5").font = plain
        sa.cell(16, 2 + g, f"={col}7+(MAX({col}8,{col}7+MAX({MIN_SPAN['body']:g},{col}11*(100-{col}7)))-{col}7)*{col}9") \
            .font = plain
        sa.cell(17, 2 + g, f"=MIN(100,100*MAX(0,({col}4-{col}3)/({col}15-{col}3))^{col}6)").font = plain
        sa.cell(18, 2 + g, f"=MIN(100,100*MAX(0,({col}8-{col}7)/({col}16-{col}7))^{col}10)").font = plain
        for r in (15, 16):
            sa.cell(r, 2 + g).number_format = "0.00"
        for r in (17, 18):
            sa.cell(r, 2 + g).number_format = "0.0"

        suffix = f" {label}" if label else ""
        raw = wb.create_sheet(f"Raw probe scores{suffix}"[:31])
        ws = wb.create_sheet(f"Updated scores{suffix}"[:31])
        score_sheets.append(ws)
        raw_sheets.append(raw)
        rname = raw.title.replace("'", "''")
        raw.append(["Step", "Face (raw %)", "Body (raw %)"])
        for c in raw[1]:
            c.font, c.fill = head, head_fill
        for pr in rows:
            raw.append([pr["step"], pr.get("face"), pr.get("body")])
        for r in raw.iter_rows(min_row=2):
            r[0].font = plain
            for c in r[1:]:
                c.font, c.number_format = blue, "0.00"
        for letter, w in zip("ABC", (8, 14, 14)):
            raw.column_dimensions[letter].width = w

        ws["A1"] = f"{cfg.job_name}{' - ' + label if label else ''}: likeness on the normalized scale"
        ws["A1"].font = Font(name=FONT, bold=True, size=12)
        ws["A2"] = ("0 = a stranger (the step-0 probe), 100 = the scale's top; where your references' own level "
                    "lands is on the Scale sheet. Smoothed = what the Monitor and the best checkpoint use. Edit the "
                    "Scale sheet's blue cells to rescore.")
        ws["A2"].font = note
        heads = ["Step", "Face (smoothed)", "Body (smoothed)", "Combined (smoothed)", "Face (probe)", "Body (probe)",
                 "Combined (probe)", "", "Face raw smoothed", "Body raw smoothed"]
        for j, h in enumerate(heads, 1):
            if h:
                c = ws.cell(4, j, h)
                c.font, c.fill, c.alignment = head, head_fill, Alignment(horizontal="center", wrap_text=True)
        s = lambda r: f"Scale!${col}${r}"  # noqa: E731

        def nrm(x, floor_row, top_row, curve_row):
            return (f"MIN(100,100*MAX(0,({x}-{s(floor_row)})/({s(top_row)}-{s(floor_row)}))^{s(curve_row)})")

        for i in range(len(rows)):
            r, rr = 5 + i, 2 + i
            face, body = f"'{rname}'!B{rr}", f"'{rname}'!C{rr}"
            ws.cell(r, 1, f"='{rname}'!A{rr}").font = green
            if i == 0:
                ws.cell(r, 9, f"={face}")
                ws.cell(r, 10, f"={body}")
            else:  # the controller's smoothing, skipping a probe without that score
                ws.cell(r, 9, f'=IF({face}="",I{r - 1},{s(14)}*I{r - 1}+(1-{s(14)})*{face})')
                ws.cell(r, 10, f'=IF({body}="",J{r - 1},{s(14)}*J{r - 1}+(1-{s(14)})*{body})')
            ws.cell(r, 2, f'=IF(I{r}="","",{nrm(f"I{r}", 3, 15, 6)})')
            ws.cell(r, 3, f'=IF(J{r}="","",{nrm(f"J{r}", 7, 16, 10)})')
            ws.cell(r, 5, f'=IF({face}="","",{nrm(face, 3, 15, 6)})')
            ws.cell(r, 6, f'=IF({body}="","",{nrm(body, 7, 16, 10)})')
            for comb, a, b in ((4, "B", "C"), (7, "E", "F")):
                ws.cell(r, comb, f'=IF(AND({a}{r}="",{b}{r}=""),"",(IF({a}{r}="",0,{a}{r}*{s(12)})'
                                 f'+IF({b}{r}="",0,{b}{r}*{s(13)}))/(IF({a}{r}="",0,{s(12)})'
                                 f'+IF({b}{r}="",0,{s(13)})))')
            for j in range(1, 11):
                c = ws.cell(r, j)
                if j > 1:
                    c.font, c.number_format = plain, "0.0"
                c.border = line
        last = 4 + len(rows)
        ws.cell(last + 2, 1, "Best smoothed combined").font = bold
        ws.cell(last + 2, 4, f"=MAX(D5:D{last})").font = bold
        ws.cell(last + 2, 4).number_format = "0.0"
        ws.cell(last + 3, 1, "at step").font = bold
        ws.cell(last + 3, 4, f"=INDEX(A5:A{last},MATCH(D{last + 2},D5:D{last},0))").font = bold
        for letter in "IJ":
            ws.column_dimensions[letter].hidden = True
        ws.column_dimensions["A"].width = 10
        for letter in "BCDEFG":
            ws.column_dimensions[letter].width = 14
        ws.row_dimensions[4].height = 32
        ws.freeze_panes = "B5"

    wb._sheets = [*score_sheets, *raw_sheets, sa]  # the scores first, the inputs after
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # openpyxl writes no cached values: the reader computes them on opening
    return wb


def config_for(values: dict):
    """A TrainConfig from saved values (a history record's config): the defaults for what it doesn't hold, keys a newer
    or older version doesn't know left out."""
    import dataclasses

    from .config import TrainConfig

    known = {f.name for f in dataclasses.fields(TrainConfig)}
    cfg = TrainConfig()
    for k, v in (values or {}).items():
        if k in known:
            setattr(cfg, k, v)
    return cfg


def export_run(record: dict, out: Path) -> Path:
    """Write a Previous Runs record's spreadsheet (its probe_raw and config) to `out`. Raises FileNotFoundError for a
    record without raw probe scores (saved before the history kept them)."""
    raw = record.get("probe_raw") or {}
    if not raw.get("records"):
        raise FileNotFoundError("This run was saved before the history kept its raw probe scores")
    cfg = config_for({**(record.get("config") or {}), "job_name": record.get("job") or "run"})
    build_workbook_from(raw["records"], cfg, raw.get("names") or []).save(out)
    return Path(out)


def export(job_dir: Path, out: Path | None = None) -> Path:
    """Write <job>_likeness.xlsx into the job folder (or `out`); returns its path."""
    job_dir = Path(job_dir)
    out = Path(out) if out else job_dir / f"{job_dir.name}_likeness.xlsx"
    build_workbook(job_dir).save(out)
    return out
