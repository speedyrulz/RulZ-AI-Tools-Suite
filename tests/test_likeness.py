import json
import random
import sys
from pathlib import Path

import pytest
import torch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from rulztrainer import likeness as L  # noqa: E402
from rulztrainer.config import TrainConfig  # noqa: E402
from rulztrainer.dataset import BucketBatcher, Item  # noqa: E402


def _cfg(**kw):
    """Controller settings; the mechanics tests run on the raw scores (likeness_normalize off), the normalization
    tests switch it on."""
    base = dict(probe_every=10, probe_ema=0.0, likeness_noise_band=1.0, likeness_track_best=True,
                likeness_patience=3, likeness_lr_control=True, likeness_lr_min=0.5, likeness_lr_max=2.0,
                likeness_lr_step=0.1, likeness_curriculum=True, likeness_curriculum_strength=1.0,
                likeness_normalize=False, face_scale_span=1.0, face_scale_curve=1.0,
                body_scale_span=1.0, body_scale_curve=1.0)  # the plain linear scales
    return TrainConfig(**{**base, **kw})


def test_controller_best_stop_lr_and_curriculum():
    c = L.LikenessController(_cfg())
    # rising scores: best updates, lr factor climbs
    rec, act = c.update(0, 40.0, 80.0)
    assert act.get("save_best") and rec.best == rec.combined_ema
    rec, act = c.update(10, 50.0, 80.0)
    assert act.get("save_best") and act.get("lr_ratio", 0) > 1
    assert c.lr_factor > 1.0
    # face lags body -> positive face bias
    assert c.face_bias > 0 and act.get("face_bias", c.face_bias) > 0
    # falling scores: lr factor shrinks, patience counts, then stop
    stops = []
    for i, f in enumerate([45.0, 40.0, 35.0]):
        rec, act = c.update(20 + 10 * i, f, 80.0)
        stops.append(bool(act.get("stop")))
    assert c.lr_factor < 1.1 and stops[-1] and c.stopped
    assert c.best_step == 10
    # bounds
    for _ in range(30):
        c.update(100, 90.0, 90.0)
    assert c.lr_factor <= 2.0


def test_controller_handles_missing_modalities():
    c = L.LikenessController(_cfg())
    rec, _ = c.update(0, None, 70.0)
    assert rec.combined_ema == 70.0 and c.face_bias == 0.0
    rec, _ = c.update(10, 60.0, None)
    assert rec.face_ema == 60.0 and rec.body_ema == 70.0


def test_apply_lr_ratio_adamw_with_scheduler_and_automagic():
    from rulztrainer.trainer import build_optimizer, build_scheduler, current_lr

    p = torch.nn.Parameter(torch.zeros(4, 4))
    cfg = TrainConfig(optimizer="adamw", learning_rate=1e-4, lr_scheduler="constant", warmup_steps=0)
    opt = build_optimizer(cfg, [p])
    sch = build_scheduler(cfg, opt)
    L.apply_lr_ratio(opt, sch, 1.5)
    p.grad = torch.zeros_like(p)
    opt.step()
    sch.step()
    assert abs(current_lr(opt, sch) - 1.5e-4) < 1e-12  # survives the scheduler step
    cfg2 = TrainConfig(optimizer="automagic3", learning_rate=1e-4)
    p2 = torch.nn.Parameter(torch.zeros(4, 4))
    opt2 = build_optimizer(cfg2, [p2])
    sch2 = build_scheduler(cfg2, opt2)
    p2.grad = torch.randn_like(p2)
    opt2.step()
    before = current_lr(opt2, sch2)
    L.apply_lr_ratio(opt2, sch2, 0.5)
    assert abs(current_lr(opt2, sch2) - before * 0.5) < 1e-9


def test_batcher_weights_and_curriculum(tmp_path):
    items = [Item(image=tmp_path / f"{i}.png", caption="", width=512, height=512, latent_path=tmp_path / "l",
                  latent_flip_path=None, text_path=tmp_path / "t") for i in range(10)]
    fractions = {str(it.image.resolve()): (0.1 if i < 3 else 0.0) for i, it in enumerate(items)}
    w = L.curriculum_weights(items, fractions, face_bias=1.0)
    assert all(w[str(it.image)] == 2.0 for it in items[:3]) and all(w[str(it.image)] == 0.2 for it in items[3:])
    b = BucketBatcher(items, 1, 0)
    b.set_weights(w)
    picks = [batch[0] for _ in range(20) for batch in b.epoch()]
    face_share = sum(1 for it in picks if str(it.image) in {str(x.image) for x in items[:3]}) / len(picks)
    assert face_share > 0.6  # 3 of 10 images get most of the draws
    b.set_weights(None)
    picks = [batch[0] for batch in b.epoch()]
    assert len(set(str(it.image) for it in picks)) == 10  # uniform shuffle visits every image
    assert L.timestep_logit_shift(1.0) < 0 < L.timestep_logit_shift(-1.0)
    assert L.probe_body_size(512) == (416, 624) and L.probe_body_size(1024) == (832, 1248)  # ~ the face's area
    assert L.probe_body_size(256) == (256, 384)


def test_normalized_scores_put_face_and_body_on_one_scale():
    """Raw, a body score moves ~10 points over a run and a face score ~60. Normalized (0 = the step-0 probe, 100 = the
    references' self-similarity) both count as the weights say, the curriculum reads the gap on that scale, and the
    log / record keep the raw values next to the normalized ones."""
    c = L.LikenessController(_cfg(likeness_normalize=True, face_weight=0.5, body_weight=0.5), face_ceiling=70.0,
                             body_ceiling=94.0)
    rec, _ = c.update(0, 15.0, 82.0)  # the base model: the floors
    assert c.scale("face") == (15.0, 70.0) and c.scale("body") == (82.0, 94.0)
    assert rec.face_norm == 0.0 and rec.body_norm == 0.0 and rec.combined_ema == 0.0
    assert rec.face == 15.0 and rec.body == 82.0 and rec.scale == {"face": [15.0, 70.0], "body": [82.0, 94.0], "curve": {"face": 1.0, "body": 1.0},
                                                                  "span": {"face": 1.0, "body": 1.0}}
    rec, act = c.update(10, 42.5, 88.0)  # both halfway to the references
    assert abs(rec.face_norm - 50.0) < 1e-9 and abs(rec.body_norm - 50.0) < 1e-9 and abs(rec.combined_ema - 50.0) < 1e-9
    assert abs(c.face_bias) < 1e-9  # no gap on the normalized scale (raw: 88 - 42.5 = 45.5 points)
    rec, act = c.update(20, 42.5, 94.0)  # the body reaches the references, the face does not: face-heavy
    assert rec.body_norm == 100.0 and c.face_bias > 0.9
    rec, _ = c.update(30, 80.0, 82.0)  # above a typical reference stops at 100; the floor clips at 0
    assert rec.face_norm == 100.0
    assert rec.body_norm == 0.0


def test_the_face_scale_is_stretched_and_curved_so_it_keeps_room_to_improve():
    """An ArcFace cosine saturates: on a real run a decent face read 90 while the body, on the same pictures, read 51.
    The face's 100 sits face_scale_span x as far above a stranger as the references do, and the position on that scale
    is raised to face_scale_curve; the body likewise (body_scale_span / _curve), so both read their references'
    own level alike."""
    plain = L.LikenessController(_cfg(likeness_normalize=True), face_ceiling=73.7, body_ceiling=88.4)
    new = L.LikenessController(TrainConfig(likeness_normalize=True, probe_ema=0.0), face_ceiling=73.7,
                               body_ceiling=88.4)
    assert (new.cfg.face_scale_span, new.cfg.face_scale_curve) == (L.FACE_SPAN_DEFAULT, L.FACE_CURVE_DEFAULT)
    for c in (plain, new):
        c.update(0, 13.8, 73.3)  # the floors
    raw_face = 13.8 + 0.9 * (73.7 - 13.8)  # 90 on the old scale
    assert abs(plain.norm("face", raw_face) - 90.0) < 1e-9
    face = new.norm("face", raw_face)
    assert 50 < face < 60 and abs(face - 54.4) < 0.5  # about what the body read for the same likeness
    assert new.norm("body", 81.0) < plain.norm("body", 81.0)  # the body is stretched and curved too
    ref_level = new.norm("face", 73.7)  # the references' own face level: no longer the ceiling
    assert 60 < ref_level < 68 and new.norm("face", 80.0) > ref_level  # and a better likeness still counts
    assert new.norm("face", 13.8) == 0.0 and new.norm("face", 200.0) == 100.0
    lo, hi = new.scale("face")
    assert lo == 13.8 and abs(hi - (13.8 + (73.7 - 13.8) * L.FACE_SPAN_DEFAULT)) < 1e-9
    assert new.scale_record()["curve"] == {"face": L.FACE_CURVE_DEFAULT, "body": L.BODY_CURVE_DEFAULT}
    assert "curve x1.5" in new.describe_scale()
    assert L.normalize(50.0, 0.0, 100.0, 2.0) == 25.0 and L.normalize(150.0, 0.0, 100.0) == 100.0


def test_normalization_defaults_floor_clamp_and_switch():
    c = L.LikenessController(_cfg(likeness_normalize=True))  # no step-0 probe, no reference self-similarity
    rec, _ = c.update(25, 20.0, 80.0)
    assert c.scale("face") == (20.0, 75.0) and c.scale("body") == (80.0, 90.0)
    assert rec.face_norm == 0.0 and rec.body_norm == 0.0 and c.describe_scale().count("default") == 2
    # the step-0 probe above the measured ceiling (a base model that already resembles the references, or a body
    # self-similarity below what probes score): step 0 is still 0, the ceiling moves up to floor + MIN_SPAN
    c = L.LikenessController(_cfg(likeness_normalize=True), face_ceiling=60.0, body_ceiling=85.0)
    rec, _ = c.update(0, 58.0, 90.0)
    assert c.scale("face") == (58.0, 68.0) and c.scale("body") == (90.0, 98.0)
    assert rec.face_norm == 0.0 and rec.body_norm == 0.0 and rec.combined_ema == 0.0
    rec, _ = c.update(10, 63.0, 94.0)
    assert abs(rec.face_norm - 50.0) < 1e-9 and abs(rec.body_norm - 50.0) < 1e-9
    assert c.describe_scale().count("step-0 probe") == 2
    # the switch off: the old raw behaviour
    c = L.LikenessController(_cfg())
    rec, _ = c.update(0, 15.0, 82.0)
    assert rec.combined_ema == 0.7 * 15.0 + 0.3 * 82.0 and rec.face_norm == 15.0 and rec.scale == {}
    assert "face_norm_ema" in c.state() and c.state()["normalize"] is False


def test_caption_dropout_keeps_the_trigger_word(tmp_path):
    """Caption dropout trains on the trigger word alone (the way the LoRA is prompted), per family's trigger
    placement; the switch off gives the empty prompt (FIBO: its null caption)."""
    from rulztrainer import dataset as D

    cfg = TrainConfig(trigger_word="zkrea", cache_dir=str(tmp_path), model_arch="krea2")
    assert D.dropout_prompt(cfg) == "zkrea"
    assert D.dropout_prompt(TrainConfig(trigger_word="", cache_dir=str(tmp_path), model_arch="krea2")) == ""
    assert D.dropout_prompt(TrainConfig(trigger_word="zkrea", caption_dropout_keep_trigger=False,
                                        cache_dir=str(tmp_path), model_arch="krea2")) == ""
    fibo = TrainConfig(trigger_word="zkrea", cache_dir=str(tmp_path), model_arch="fibo2")
    assert '"short_description":"zkrea"' in D.dropout_prompt(fibo)
    fibo_off = TrainConfig(trigger_word="zkrea", caption_dropout_keep_trigger=False, cache_dir=str(tmp_path),
                           model_arch="fibo2")
    assert '"short_description":""' in D.dropout_prompt(fibo_off)
    assert D.dropout_text_path(cfg) != D.dropout_text_path(fibo_off)


def test_sample_scores_carry_normalized_values_and_captions_show_them(tmp_path):
    """Preview samples are scored on the probes' normalized scale too (score_samples(scale=controller)), and the
    caption shows the normalized number with the raw one behind it; without a scale, raw only."""
    from types import SimpleNamespace

    from rulztrainer import facesim

    c = L.LikenessController(_cfg(likeness_normalize=True, face_weight=0.5, body_weight=0.5), face_ceiling=70.0,
                             body_ceiling=60.0)
    c.update(0, 15.0, 8.0)  # the floors
    face = SimpleNamespace(score=lambda p, step, i: facesim.FaceScore(step, str(p), i, 1, 0.425, 42.5, "ref"))
    body = SimpleNamespace(score=lambda p: (34.0, 3))
    (tmp_path / "s.png").write_bytes(b"")
    scores = facesim.score_samples(face, [str(tmp_path / "s.png")], 10, tmp_path, body, 0.5, 0.5, scale=c)
    s = scores[0]
    assert s.percent == 42.5 and s.body_percent == 34.0 and abs(s.combined_percent - 38.25) < 1e-9
    assert abs(s.face_norm - 50.0) < 1e-9 and abs(s.body_norm - 50.0) < 1e-9 and abs(s.combined_norm - 50.0) < 1e-9
    import webui

    assert webui.score_caption(facesim.asdict(s)) == "  |  face 50%, body 50%, combined 50%"
    raw = facesim.score_samples(face, [str(tmp_path / "s.png")], 10, tmp_path, body, 0.5, 0.5)[0]
    assert raw.face_norm is None and webui.score_caption(facesim.asdict(raw)) == "  |  face 42%, body 34%, combined 38%"
    assert len(facesim.read_scores(tmp_path)) == 2


def test_reference_ceiling_is_the_median_best_match():
    """Each reference's best match among the others, the median over the usable ones: one odd photo does not drag
    the 100% down the way it drags a mean."""
    vecs = [[1.0, 0.0, 0.0], [0.8, 0.6, 0.0], [0.0, 1.0, 0.0], [0.6, 0.8, 0.0], [0.0, 0.0, -1.0]]
    assert abs(L.reference_ceiling(vecs) - 0.8) < 1e-6  # best matches .8 .96 .8 .96 0 (mean .70)
    assert abs(L.reference_ceiling(vecs, [True, True, True, True, False]) - 0.88) < 1e-6  # the odd one left out
    assert abs(L.reference_ceiling(vecs, [True, False, False, False, False]) - 0.8) < 1e-6  # < 2 usable: all
    assert L.reference_ceiling(vecs[:1]) is None and L.reference_ceiling([]) is None


def test_the_step0_probe_renders_more_seeds():
    cfg = TrainConfig(probe_seeds=2, probe_floor_seeds=4)
    assert L.probe_seed_count(cfg, 0) == 4 and L.probe_seed_count(cfg, 60) == 2
    assert L.probe_seed_count(TrainConfig(probe_seeds=6, probe_floor_seeds=4), 0) == 6  # never fewer than usual
    assert L.probe_seed_count(TrainConfig(probe_seeds=2, probe_floor_seeds=4, likeness_normalize=False), 0) == 2
    assert L.probe_seed_count(TrainConfig(probe_seeds=0, probe_floor_seeds=0), 0) == 1


def test_the_body_scale_is_stretched_like_the_face_and_never_too_narrow():
    """A reference-level body reads like a reference-level face (~64), and a body ceiling barely above the step-0 score
    is widened to halfway to 100 (a real run's 73.3 -> 81.3 read 100 at a few points of copied outfit)."""
    from types import SimpleNamespace

    cfg = SimpleNamespace(likeness_normalize=True, probe_ema=0.5, resolution_step_up=False)
    c = L.LikenessController(cfg, face_ceiling=70.0, body_ceiling=88.0)
    c.floor.update(face=15.0, body=30.0)
    assert abs(c.norm("body", 88.0) - c.norm("face", 70.0)) < 1e-9 and 60 < c.norm("body", 88.0) < 66
    assert c.norm("body", 100.0) < 90  # a probe that copies the photos no longer reaches 100
    narrow = L.LikenessController(cfg, body_ceiling=81.3)
    narrow.floor["body"] = 73.3
    lo, hi = narrow.scale("body")
    assert abs(hi - (73.3 + 0.5 * (100 - 73.3) * L.BODY_SPAN_DEFAULT)) < 1e-9
    assert narrow.norm("body", 81.3) < 35
    linear = L.LikenessController(SimpleNamespace(**vars(cfg), body_scale_span=1.0, body_scale_curve=1.0),
                                  body_ceiling=88.0)
    linear.floor["body"] = 30.0
    assert linear.norm("body", 88.0) == 100.0  # 1 and 1: the old body scale


def test_likeness_spreadsheet_export(tmp_path):
    """The Monitor's Export spreadsheet: one row per probe (the step-0 probe logged twice kept once), the run's floors and
    references' own levels on the Scale sheet (an old record's stretch taken back out), live formulas, and a race's
    configurations each on their own sheets."""
    from openpyxl import load_workbook

    from rulztrainer import likeness_export as X

    old_scale = {"face": [11.0, 11.0 + 60.0 * 1.35], "body": [26.0, 88.0], "curve": {"face": 1.5, "body": 1.0}}

    def rec(step, face, body, prefix="", scale=old_scale):
        return {"step": step, "face": face, "body": body, "scale": scale,
                "images": [f"probes/{prefix}step{step:06d}_face_0.png"]}

    job = tmp_path / "run"
    job.mkdir()
    (job / "likeness.jsonl").write_text("\n".join(json.dumps(r) for r in (
        rec(0, 11.0, 26.0), rec(0, 11.0, 26.0), rec(100, 40.0, 60.0), rec(200, 70.0, None))))
    out = X.export(job)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Updated scores", "Raw probe scores", "Scale"]
    sc, ws, raw = wb["Scale"], wb["Updated scores"], wb["Raw probe scores"]
    assert abs(sc["B4"].value - 71.0) < 1e-9 and sc["B8"].value == 88.0  # the references' own levels, unstretched
    assert sc["B5"].value == 1.35 and sc["B9"].value == L.BODY_SPAN_DEFAULT  # rescored on the current body scale
    assert [raw.cell(r, 1).value for r in range(2, 5)] == [0, 100, 200] and raw["C4"].value is None
    assert ws["B6"].value.startswith('=IF(I6="","",MIN(100,100*MAX(0,(I6-Scale!$B$3)')
    assert ws["D9"].value == "=MAX(D5:D7)" and "INDEX(A5:A7" in ws["D10"].value
    assert wb.calculation.fullCalcOnLoad

    race = tmp_path / "race"
    race.mkdir()
    (race / "race.json").write_text(json.dumps({"names": ["lora16", "lokr8"]}))
    new_scale = {"face": [10.0, 91.0], "body": [30.0, 108.0], "curve": {"face": 1.5, "body": 1.5},
                 "span": {"face": 1.35, "body": 1.3}}
    (race / "likeness.jsonl").write_text("\n".join(json.dumps(r) for r in (
        rec(0, 10.0, 30.0, "lora16_", new_scale), rec(0, 12.0, 31.0, "lokr8_", new_scale),
        rec(100, 50.0, 70.0, "lokr8_", new_scale), rec(300, 60.0, 80.0, "", new_scale))))
    wb = X.build_workbook(race)
    assert wb.sheetnames == ["Updated scores lora16", "Updated scores lokr8", "Updated scores winner",
                             "Raw probe scores lora16", "Raw probe scores lokr8", "Raw probe scores winner", "Scale"]
    sc = wb["Scale"]
    assert [sc.cell(2, c).value for c in (2, 3, 4)] == ["lora16", "lokr8", "winner"]
    assert abs(sc["C8"].value - (30.0 + 78.0 / 1.3)) < 1e-9  # the logged span taken out
    assert "Scale!$C$3" in wb["Updated scores lokr8"]["B5"].value
    assert wb["Raw probe scores lokr8"].max_row == 3
    with pytest.raises(FileNotFoundError):
        X.build_workbook(tmp_path / "empty")
