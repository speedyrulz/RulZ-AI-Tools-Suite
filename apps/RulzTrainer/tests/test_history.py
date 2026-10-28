"""Run history (output/_history): the series kept from a run's scores, saving while it trains, labels, removal, importing
earlier job folders, and the Previous Runs table and overlay series."""
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import webui  # noqa: E402
from rulztrainer import history  # noqa: E402
from rulztrainer.config import TrainConfig  # noqa: E402

SCORES = [{"step": 300, "face_norm": 10.0, "body_norm": 0.0, "combined_norm": 5.0, "percent": 20.0},
          {"step": 300, "face_norm": 20.0, "body_norm": None, "combined_norm": 15.0, "percent": 30.0},
          {"step": 600, "face_norm": 40.0, "body_norm": 30.0, "combined_norm": 35.0, "percent": 50.0}]
PROBES = [{"step": 0, "face_norm_ema": 0.0, "body_norm_ema": 0.0, "combined_ema": 0.0},
          {"step": 25, "face_norm_ema": 12.34, "body_norm_ema": 5.0, "combined_ema": 9.0},
          {"step": 50, "face_ema": 60.0, "body_ema": 55.0, "combined_ema": 58.0}]  # an older record: raw only


def _status(**kw):
    return {"step": 600, "total": 2000, "face_scores": SCORES, "probes": PROBES, **kw}


def test_series_are_means_per_step_and_fall_back_to_raw_scores():
    ident = history.identity_series(SCORES)
    assert ident["normalized"] and ident["face"] == [[300, 15.0], [600, 40.0]] and ident["body"] == [[300, 0.0], [600, 30.0]]
    raw = history.identity_series([{"step": 1, "percent": 30.0, "body_percent": 50.0}])
    assert not raw["normalized"] and raw["face"] == [[1, 30.0]] and raw["body"] == [[1, 50.0]]
    probes = history.probe_series(PROBES)
    assert probes["face"] == [[0, 0.0], [25, 12.3], [50, 60.0]] and probes["combined"][-1] == [50, 58.0]
    assert history.probe_series([]) == {"face": [], "body": [], "combined": []}


def test_a_run_is_saved_while_it_trains_keeps_its_label_and_needs_data(tmp_path):
    cfg = TrainConfig(job_name="my lora", lora_rank=32, steps=2000)
    rid = history.new_run_id(cfg.job_name, 1_700_000_000)
    assert rid.startswith("my_lora-2023")
    assert history.save_run(cfg, rid, 1_700_000_000, {"step": 5, "face_scores": [], "probes": []}, root=tmp_path) is None
    path = history.save_run(cfg, rid, 1_700_000_000, _status(), root=tmp_path)
    assert path == tmp_path / f"{rid}.json"
    assert history.set_label(rid, "  rank 32 test ", root=tmp_path) and not history.set_label("nope", "x", root=tmp_path)
    history.save_run(cfg, rid, 1_700_000_000, _status(step=900), "finished", root=tmp_path)  # a later save
    (run,) = history.list_runs(tmp_path)
    assert run["label"] == "rank 32 test" and run["state"] == "finished" and run["step"] == 900
    assert run["config"]["lora_rank"] == 32 and run["total"] == 2000 and run["arch"] == "moss_image"
    assert history.delete_runs([rid, "../x"], tmp_path) == 1 and history.list_runs(tmp_path) == []


def test_signature_changes_when_the_history_does(tmp_path):
    cfg = TrainConfig(job_name="a")
    empty = history.signature(tmp_path)
    history.save_run(cfg, "a-1", 1.0, _status(), root=tmp_path)
    first = history.signature(tmp_path)
    assert first != empty
    history.set_label("a-1", "x", tmp_path)
    assert history.signature(tmp_path) != first


def _job(output: Path, name: str, **status):
    job = output / "group" / name
    job.mkdir(parents=True)
    (job / "config.yaml").write_text(f"job_name: {name}\noutput_dir: {output / 'group'}\nsteps: 2000\n"
                                     "network_type: lokr\n", encoding="utf-8")
    (job / "status.json").write_text(json.dumps({"step": 2000, "total": 2000, "done": True, **status}),
                                     encoding="utf-8")
    return job


def test_earlier_job_folders_are_imported_once(tmp_path):
    out, root = tmp_path / "output", tmp_path / "output" / "_history"
    _job(out, "scored", face_scores=SCORES, probes=PROBES)
    _job(out, "unscored")
    _job(out, "_crops_like_job_without_scores")
    assert history.import_jobs(out, root) == 1
    assert history.import_jobs(out, root) == 0  # the job already has a run
    (run,) = history.list_runs(root)
    assert run["job"] == "scored" and run["state"] == "finished" and run["id"].endswith("-imported")
    assert run["config"]["network_type"] == "lokr" and run["probes"]["combined"]


def test_the_table_and_the_overlay_series():
    runs = [{"id": "a", "job": "krea", "started": 1_700_000_000, "label": "", "state": "finished", "step": 2000,
             "total": 2000, "config": {"network_type": "lokr", "lokr_factor": 8, "learning_rate": 1e-4},
             "identity": {"face": [[300, 15.0]], "body": [[300, 5.0]]},
             "probes": {"face": [[25, 40.0], [50, 60.0]], "body": [[25, 10.0]], "combined": [[25, 20.0], [50, 70.0]]}},
            {"id": "b", "job": "moss", "started": 1_700_100_000, "label": "rank 16", "state": "running", "step": 10,
             "total": 100, "config": {"lora_rank": 16, "finetune": False},
             "identity": {"face": [], "body": []}, "probes": {"face": [], "body": [], "combined": []}}]
    rows = webui.history_rows(runs)
    assert rows[0]["adapter"] == "LoKr f8" and rows[0]["steps"] == "2,000 / 2,000" and rows[0]["lr"] == "1.0e-04"
    assert (rows[0]["face"], rows[0]["body"], rows[0]["combined"], rows[0]["best"]) == ("60.0%", "10.0%", "70.0%", "70.0%")
    assert rows[1]["run"] == "rank 16" and rows[1]["adapter"] == "LoRA r16" and rows[1]["best"] == "-"
    probes = webui.history_series(runs, "probes", ["face", "body", "combined"])
    assert [s["name"].rsplit(" - ", 1)[1] for s in probes] == ["face", "body", "combined"]  # run b has no points
    assert probes[1]["lineStyle"]["type"] == "dashed" and probes[2]["lineStyle"]["width"] == 3
    assert all(s["itemStyle"]["color"] == webui.HISTORY_COLORS[0] for s in probes)
    both = webui.history_series(runs, "identity", ["face"])
    assert len(both) == 1 and both[0]["data"] == [[300, 15.0]] and both[0]["showSymbol"]
    assert webui.history_series(runs, "probes", []) == []


def test_a_race_keeps_each_configuration_s_probes_apart(tmp_path):
    def rec(name, step, comb):
        prefix = f"{name}_" if name else ""
        return {"step": step, "face_norm_ema": comb - 5, "body_norm_ema": comb + 5, "combined_ema": comb,
                "images": [f"out/probes/{prefix}step{step:06d}_face_0.png"]}

    status = {"race": {"names": ["lora16", "lokr8"]},
              "probes": [rec("lora16", 50, 10.0), rec("lokr8", 50, 14.0), rec("lokr8", 100, 20.0), rec("", 150, 25.0)]}
    series = history.probe_series(status["probes"], history.race_names(status))
    assert series["combined"] == [[50, 10.0], [50, 14.0], [100, 20.0], [150, 25.0]]  # all together, as before
    assert list(series["racers"]) == ["lora16", "lokr8", "winner"]
    assert series["racers"]["lokr8"]["combined"] == [[50, 14.0], [100, 20.0]]
    assert series["racers"]["winner"]["face"] == [[150, 20.0]]
    run = {"id": "r", "job": "race", "label": "my race", "probes": series, "identity": {}, "config": {"race": True}}
    lines = webui.history_series([run], "probes", ["combined"])
    assert [s["name"] for s in lines] == ["my race lora16 - combined", "my race lokr8 - combined",
                                          "my race winner - combined"]
    assert len({s["itemStyle"]["color"] for s in lines}) == 3  # one colour per configuration
    # a race saved before the split gets it from its job folder on import
    job = tmp_path / "output" / "race"
    job.mkdir(parents=True)
    (job / "status.json").write_text(json.dumps(status))
    (job / "config.yaml").write_text("job_name: race\nrace: true\n")
    root = tmp_path / "history"
    old = {**run, "probes": {k: v for k, v in series.items() if k != "racers"}}
    history._write(root / "r.json", old)
    assert history.import_jobs(tmp_path / "output", root) == 1
    upgraded = history.list_runs(root)
    assert len(upgraded) == 1 and list(upgraded[0]["probes"]["racers"]) == ["lora16", "lokr8", "winner"]
    assert upgraded[0]["label"] == "my race"


def test_probe_images_are_kept_with_their_run_and_compared_at_the_same_step(tmp_path):
    from PIL import Image

    job = tmp_path / "out" / "probes"
    job.mkdir(parents=True)
    for name in ("lora16_step000100_face_0.png", "lora16_step000100_face_1.png", "lokr8_step000100_face_0.png",
                 "lokr8_step000200_face_0.png", "step000300_face_0.png"):
        Image.new("RGB", (1024, 1024), "white").save(job / name)

    def rec(step, prefix, seeds, comb):
        return {"step": step, "face_norm": comb, "body_norm": None, "combined": comb,
                "images": [str(job / f"{prefix}step{step:06d}_face_{k}.png") for k in seeds]}

    status = {"race": {"names": ["lora16", "lokr8"]}, "probes": [
        rec(100, "lora16_", (0, 1), 30.0), rec(100, "lokr8_", (0,), 35.0), rec(200, "lokr8_", (0,), 50.0),
        rec(300, "", (0,), 60.0)]}
    root = tmp_path / "history"
    kept = history.keep_probe_images("run1", status, root)
    assert len(kept) == 5 and all((root / e["file"]).is_file() for e in kept)
    assert {(e["racer"], e["step"], e["seed"]) for e in kept} >= {("lora16", 100, 1), ("", 300, 0)}
    with Image.open(root / kept[0]["file"]) as im:
        assert max(im.size) == history.PROBE_THUMB_PX and im.format == "JPEG"
    assert next(e for e in kept if e["racer"] == "lokr8" and e["step"] == 200)["combined"] == 50.0
    (job / "lora16_step000100_face_0.png").unlink()  # copies outlive the job folder's own (a reused job name)
    assert len(history.keep_probe_images("run1", status, root)) == 5

    run = {"id": "run1", "job": "race", "label": "my race", "probe_images": kept, "config": {"probe_size": 512}}
    groups = webui.probe_image_groups([run])
    assert [g[0] for g in groups] == ["my race lora16", "my race lokr8", "my race winner"]
    assert webui.probe_image_steps(groups, "face") == [100, 200, 300] and webui.probe_image_steps(groups, "body") == []
    at, imgs = webui.probe_images_at(groups[0][2], 200, "face")  # lora16 left the race after step 100
    assert at == 100 and [e["seed"] for e in imgs] == [0, 1]
    assert webui.probe_images_at(groups[2][2], 200, "face") == (None, [])  # the winner only from step 300
    other = {"id": "run2", "job": "x", "label": "", "probe_images": [], "config": {"probe_size": 768}}
    assert webui.probe_settings_differences([run, other]) == ["probe_size: 512 / 768"]
    assert webui.probe_caption({"combined": 61.4, "face": 70.0, "body": None}) == "overall 61 | face 70 | body -"
    assert history.delete_runs(["run1"], root) == 0 and not (root / "run1").exists()  # the copies go too


def _raw_probe(step, face, body, prefix="", scale=None):
    return {"step": step, "face": face, "body": body, "scale": scale or {"face": [10.0, 91.0]},
            "images": [f"output/j/probes/{prefix}step{step:06d}_face_0.png"]}


def test_raw_probe_scores_are_kept_for_the_spreadsheet(tmp_path):
    from openpyxl import load_workbook

    from rulztrainer import likeness_export

    probes = [_raw_probe(0, 10.0, 30.0), _raw_probe(0, 10.0, 30.0), _raw_probe(100, 50.0, 70.0),
              _raw_probe(200, 55.0, None, scale={"face": [10.0, 95.0]})]
    raw = history.probe_raw({"probes": probes})
    assert raw["names"] == [] and [r["step"] for r in raw["records"]] == [0, 100, 200]  # step 0 kept once
    assert "scale" in raw["records"][0] and "scale" not in raw["records"][1] and "scale" in raw["records"][2]
    race = history.probe_raw({"race": {"names": ["a", "b"]},
                              "probes": [_raw_probe(0, 10.0, 30.0, "a_"), _raw_probe(0, 11.0, 31.0, "b_")]})
    assert [(r["racer"], "scale" in r) for r in race["records"]] == [("a", True), ("b", True)]

    cfg = TrainConfig(job_name="kept", face_scale_span=1.1, face_scale_curve=3.0, face_weight=0.8, body_weight=0.2)
    history.save_run(cfg, "kept-1", 1.0, _status(probes=probes), root=tmp_path)
    (run,) = history.list_runs(tmp_path)
    assert run["probe_raw"] == raw and run["config"]["probe_ema"] == cfg.probe_ema
    out = likeness_export.export_run(run, tmp_path / "kept.xlsx")
    sc = load_workbook(out)["Scale"]
    assert (sc["B5"].value, sc["B6"].value, sc["B12"].value) == (1.1, 3.0, 0.8)  # the run's own settings
    assert load_workbook(out)["Raw probe scores"].max_row == 4
    run.pop("probe_raw")
    with pytest.raises(FileNotFoundError):
        likeness_export.export_run(run, tmp_path / "none.xlsx")


def test_import_adds_raw_scores_to_a_run_saved_before_they_were_kept(tmp_path):
    out, root = tmp_path / "output", tmp_path / "output" / "_history"
    probes = [{**_raw_probe(0, 10.0, 30.0), **PROBES[0]}, {**_raw_probe(25, 40.0, 60.0), **PROBES[1]}]
    _job(out, "old", face_scores=SCORES, probes=probes)
    assert history.import_jobs(out, root) == 1
    (run,) = history.list_runs(root)
    path = root / f"{run['id']}.json"
    run.pop("probe_raw")
    path.write_text(json.dumps(run), encoding="utf-8")  # as saved by an older version
    assert history.import_jobs(out, root) == 1
    (run,) = history.list_runs(root)
    assert [r["step"] for r in run["probe_raw"]["records"]] == [0, 25]
    assert history.import_jobs(out, root) == 0


def test_runs_compare_on_one_scale():
    scale = {"face": [10.0, 10.0 + 60.0 * 1.35], "body": [70.0, 70.0 + 15.0 * 1.35], "curve": {"face": 1.5, "body": 1.5},
             "span": {"face": 1.35, "body": 1.35}}
    probes = [{"step": 0, "face": 10.0, "body": 70.0, "scale": scale}, {"step": 100, "face": 70.0, "body": 85.0}]
    run = {"id": "r", "job": "j", "started": 0, "label": "", "config": {"probe_ema": 0.0},
           "probe_raw": {"names": [], "records": probes}, "probes": {"face": [], "body": [], "combined": []}}
    old = {**run, "id": "o", "label": "old", "probes": {"face": [[0, 0.0]], "body": [], "combined": [[0, 0.0]]}}
    old.pop("probe_raw")
    linear = {"face_scale_span": 1.0, "face_scale_curve": 1.0, "body_scale_span": 1.0, "body_scale_curve": 1.0,
              "face_weight": 0.5}
    (new, kept), own = webui.rescaled_runs([run, old], linear)
    assert new["probes"]["face"] == [[0, 0.0], [100, 100.0]] and new["probes"]["body"][-1] == [100, 100.0]
    assert new["probes"]["combined"][-1] == [100, 100.0] and kept is old and own == ["old"]
    (curved,), _ = webui.rescaled_runs([run], {**linear, "face_scale_span": 1.25, "face_scale_curve": 2.0,
                                               "face_weight": 1.0})
    assert curved["probes"]["face"][-1] == [100, 64.0] and curved["probes"]["combined"][-1] == [100, 64.0]
    assert webui.rescaled_runs([run], None) == ([run], [])

    race = {**run, "probe_raw": {"names": ["a", "b"], "records": [
        {**probes[0], "racer": "a"}, {**probes[0], "racer": "b"}, {**probes[1], "racer": "a"},
        {"step": 100, "face": 40.0, "body": 77.5, "racer": "b"}]}}
    (r,), _ = webui.rescaled_runs([race], linear)
    assert r["probes"]["racers"]["a"]["face"][-1] == [100, 100.0] and r["probes"]["racers"]["b"]["face"][-1] == [100, 50.0]
    assert [p[1] for p in r["probes"]["face"]] == [0.0, 0.0, 100.0, 50.0]  # each racer smoothed on its own
