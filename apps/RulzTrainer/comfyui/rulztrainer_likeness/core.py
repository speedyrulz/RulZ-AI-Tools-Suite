"""Likeness scoring as the RulzTrainer trainer does it, for the ComfyUI nodes (no ComfyUI imports here).

Everything that decides a score is the trainer's own code (rulztrainer.facesim, rulztrainer.bodysim,
rulztrainer.likeness): the same face model (InsightFace buffalo_l, the largest face, the best cosine against any
reference), the same body score (the default: a DINOv2 embedding of the body window, neck to mid-thigh; or the
proportions score with MediaPipe, optionally with SAM 3 silhouettes), the same normalized scale (0 = a stranger,
100 = a typical reference's best match among the others, capped at 100) and the same face / body weighting.

Averages follow the trainer's probes: the raw face scores of the images with a face are averaged, likewise the body
scores, then each average is normalized and the two are combined into the overall score. Each image also gets its
own normalized face, body and overall score."""
from __future__ import annotations

import csv
import json
import logging
import math
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from types import SimpleNamespace

from rulztrainer import bodysim, facesim
from rulztrainer.likeness import LikenessController, combine, MIN_SPAN, normalize as normalize_score

log = logging.getLogger("rulztrainer")

FACE_WEIGHT, BODY_WEIGHT = 0.7, 0.3  # the trainer's defaults (face_weight / body_weight)
_MODELS: dict[tuple, object] = {}  # loaded scorer models, shared by every node (the references live in RefSet)
_STATE = ("refs", "profile", "ref_count", "ref_self_similarity")  # a scorer's per-reference state


# ---------------------------------------------------------------- models


def face_model(provider: str = "cpu", det_size: int = 640):
    key = ("face", provider, int(det_size))
    if key not in _MODELS:
        _MODELS[key] = facesim.FaceScorer(provider, int(det_size))
    return _MODELS[key]


def body_model(kind: str, models_dir: str, device: str = "cuda:0", backend: str = "mediapipe+sam3",
               sam3_model: str = "facebook/sam3"):
    """The trainer's body scorer: appearance (the default) or proportions (MediaPipe, + SAM 3 with mediapipe+sam3,
    falling back to MediaPipe alone when SAM 3 cannot load, as the trainer does)."""
    if kind == "appearance":
        key = ("appearance", device)
        if key not in _MODELS:
            _MODELS[key] = bodysim.AppearanceScorer(models_dir, device)
        return _MODELS[key]
    key = ("proportions", backend, device, sam3_model)
    if key not in _MODELS:
        sil = None
        if backend == "mediapipe+sam3":
            try:
                sil = bodysim.SilhouetteScorer(sam3_model, device)
            except Exception as e:  # noqa: BLE001
                log.warning("SAM 3 unavailable (%s); body proportions use MediaPipe only. For facebook/sam3 accept "
                            "the license on huggingface.co and log in with 'hf auth login'.", e)
        _MODELS[key] = bodysim.BodyScorer(models_dir, silhouette=sil)
    return _MODELS[key]


def _snapshot(scorer) -> dict:
    return {k: getattr(scorer, k) for k in _STATE if hasattr(scorer, k)}


def _restore(scorer, state: dict) -> None:
    for k, v in state.items():
        setattr(scorer, k, v)


# ---------------------------------------------------------------- references


@dataclass
class RefSet:
    """Face and body references with their scorers (each kind None when it has no usable reference)."""
    face: object | None = None
    face_state: dict = field(default_factory=dict)
    face_info: list = field(default_factory=list)
    body: object | None = None
    body_state: dict = field(default_factory=dict)
    body_info: list = field(default_factory=list)
    body_kind: str = "appearance"
    labels: dict = field(default_factory=dict)  # file path -> the name shown for it (images given as a batch)
    notes: list = field(default_factory=list)

    @property
    def face_ceiling(self):
        """The references' typical best match among themselves in raw percent (None: fewer than two usable)."""
        sim = self.face_state.get("ref_self_similarity")
        return None if sim is None else float(sim) * 100.0

    @property
    def body_ceiling(self):
        sim = self.body_state.get("ref_self_similarity")
        return None if sim is None else float(sim) * 100.0

    @property
    def face_count(self) -> int:
        return len(self.face_state.get("refs") or [])

    @property
    def body_count(self) -> int:
        return int(self.body_state.get("ref_count") or 0)

    def label(self, path) -> str:
        return self.labels.get(str(path), str(path)) if path is not None else ""

    def describe(self) -> str:
        lines = []
        if self.face is not None:
            lines.append(f"Face references: {self.face_count} of {len(self.face_info)} have a face")
        else:
            lines.append("Face references: none (no face score)")
        if self.body is not None:
            what = "show the body window" if self.body_kind == "appearance" else "have a pose"
            lines.append(f"Body references ({self.body_kind}): {self.body_count} of {len(self.body_info)} {what}")
        else:
            lines.append("Body references: none (no body score)")
        return "\n".join(lines + self.notes)


def build_references(face_paths, body_paths, *, face_provider="cpu", det_size=640, body_kind="appearance",
                     body_backend="mediapipe+sam3", body_device="cuda:0", sam3_model="facebook/sam3",
                     models_dir="models", labels=None) -> RefSet:
    """Embed the references as the trainer does (set_references). face_paths / body_paths: image files or folders;
    an empty list leaves that kind out."""
    refs = RefSet(body_kind=body_kind, labels=dict(labels or {}))
    if face_paths:
        scorer = face_model(face_provider, det_size)
        refs.face_info = scorer.set_references(face_paths)
        if scorer.refs:
            refs.face, refs.face_state = scorer, _snapshot(scorer)
        else:
            refs.notes.append("No face found in any face reference: face scoring is off")
    if body_paths:
        scorer = body_model(body_kind, models_dir, body_device, body_backend, sam3_model)
        refs.body_info = scorer.set_references(body_paths)
        if scorer.ready:
            refs.body, refs.body_state = scorer, _snapshot(scorer)
        else:
            refs.notes.append(f"No usable {'person' if body_kind == 'appearance' else 'pose'} in any body reference: "
                              "body scoring is off")
    return refs


# ---------------------------------------------------------------- the scale


@dataclass
class Scale:
    """The normalized scale: per kind (floor, ceiling) in raw percent, where each came from, and whether it is on."""
    normalize: bool
    face: tuple
    body: tuple
    sources: dict
    curve: dict = field(default_factory=lambda: {"face": 1.0, "body": 1.0})

    def norm(self, kind: str, value):
        if value is None or not self.normalize:
            return value
        return normalize_score(value, *getattr(self, kind), self.curve.get(kind, 1.0))

    def describe(self) -> str:
        if not self.normalize:
            return "Normalization off: the scores are the raw similarity percentages"
        parts = []
        for k in ("face", "body"):
            lo, hi = getattr(self, k)
            parts.append(f"{k} {lo:.1f} -> {hi:.1f} raw (floor: {self.sources[k]['floor']}, "
                         f"ceiling: {self.sources[k]['ceiling']}"
                         + (f", curve x{self.curve[k]:g}" if self.curve.get(k, 1.0) != 1.0 else "") + ")")
        return "Scale: " + "; ".join(parts)


def job_scale(job_dir) -> dict:
    """The scale a trainer run used ({"face": [floor, ceiling], "body": [...]}) from its likeness.jsonl."""
    p = Path(str(job_dir).strip().strip('"'))
    p = p if p.name == "likeness.jsonl" else p / "likeness.jsonl"
    if not p.is_file():
        raise FileNotFoundError(f"No likeness.jsonl in {p.parent} (a RulzTrainer run with likeness probes)")
    scale = None
    for line in p.read_text(encoding="utf-8").splitlines():
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        if rec.get("scale"):
            scale = rec["scale"]
    if not scale:
        raise ValueError(f"{p} has no normalized scale (the run had likeness_normalize off)")
    return scale


def build_scale(refs: RefSet, normalize=True, baseline_rows=None, job=None, face_floor=None, body_floor=None,
                face_ceiling=None, body_ceiling=None) -> Scale:
    """The trainer's scale (LikenessController.scale): floor = the step-0 probe (here: the mean score of the baseline
    images, the base model without the LoRA) or the default; ceiling = the references' typical best match (2+
    references) or the default; the ceiling at least MIN_SPAN above the floor. A trainer run's own scale (job) and
    numbers typed in (>= 0) take precedence, in that order: typed > run > baseline / measured > default."""
    ctl = LikenessController(SimpleNamespace(likeness_normalize=normalize, resolution_step_up=False),
                             face_ceiling=refs.face_ceiling, body_ceiling=refs.body_ceiling)
    sources = {k: {"floor": "default", "ceiling": "references" if c is not None else "default"}
               for k, c in (("face", refs.face_ceiling), ("body", refs.body_ceiling))}
    fixed: dict[str, list] = {"face": [None, None], "body": [None, None]}  # final raw numbers, as a run logged them
    curve = dict(ctl.curve)
    for kind, key in (("face", "face_raw"), ("body", "body_raw")):
        vals = [r[key] for r in (baseline_rows or []) if r.get(key) is not None]
        if vals:
            ctl.floor[kind] = sum(vals) / len(vals)
            sources[kind]["floor"] = f"baseline images ({len(vals)})"
    if job:
        for kind in ("face", "body"):
            if kind in job:
                fixed[kind] = [float(job[kind][0]), float(job[kind][1])]
                sources[kind] = {"floor": "trainer run", "ceiling": "trainer run"}
                curve[kind] = float((job.get("curve") or {}).get(kind, 1.0))  # a run from before the curve: linear
    for kind, lo, hi in (("face", face_floor, face_ceiling), ("body", body_floor, body_ceiling)):
        if lo is not None and lo >= 0:
            fixed[kind][0] = float(lo)
            sources[kind]["floor"] = "typed in"
        if hi is not None and hi >= 0:
            fixed[kind][1] = float(hi)
            sources[kind]["ceiling"] = "typed in"
    scales = {}
    for kind in ("face", "body"):
        if fixed[kind][0] is not None:
            ctl.floor[kind] = fixed[kind][0]  # a floor typed in or logged moves the stretched span with it
        lo, hi = ctl.scale(kind)
        if fixed[kind][1] is not None:  # a ceiling typed in or logged is the final one (the trainer prints these)
            hi = max(fixed[kind][1], lo + MIN_SPAN[kind])
        scales[kind] = (lo, hi)
    return Scale(normalize, scales["face"], scales["body"], sources, curve)


# ---------------------------------------------------------------- scoring


def score_rows(refs: RefSet, items, progress=None, interrupted=None) -> list[dict]:
    """Raw scores of each (label, path): the face score as the trainer's FaceScorer.score, the body score as its body
    scorer's score. progress(i) after each image; interrupted() raises to stop."""
    if refs.face is not None:
        _restore(refs.face, refs.face_state)
    if refs.body is not None:
        _restore(refs.body, refs.body_state)
    rows = []
    for i, (label, path) in enumerate(items):
        if interrupted:
            interrupted()
        row = {"index": i + 1, "image": label, "path": str(path), "faces": None, "face_cosine": None,
               "face_raw": None, "best_face_ref": None, "body_raw": None, "body_compared": None}
        if refs.face is not None:
            s = refs.face.score(path)
            row.update(faces=s.faces, face_cosine=s.cosine, face_raw=s.percent, best_face_ref=refs.label(s.best_ref))
        if refs.body is not None:
            row["body_raw"], row["body_compared"] = refs.body.score(path)
        rows.append(row)
        if progress:
            progress(i + 1)
    return rows


def apply_scale(rows, scale: Scale, face_weight=FACE_WEIGHT, body_weight=BODY_WEIGHT) -> list[dict]:
    for r in rows:
        r["face"], r["body"] = scale.norm("face", r["face_raw"]), scale.norm("body", r["body_raw"])
        r["overall"] = combine(r["face"], r["body"], face_weight, body_weight)
        r["overall_raw"] = combine(r["face_raw"], r["body_raw"], face_weight, body_weight)
    return rows


def _mean(vals):
    return sum(vals) / len(vals) if vals else None


def summarize(rows, scale: Scale, face_weight=FACE_WEIGHT, body_weight=BODY_WEIGHT) -> dict:
    """The averages as the trainer averages a probe's seeds: the raw scores of the images that have one averaged,
    then normalized, then combined; plus the spread of the images' own normalized scores."""
    out = {"images": len(rows), "face_weight": face_weight, "body_weight": body_weight}
    for kind in ("face", "body"):
        raw = [r[f"{kind}_raw"] for r in rows if r[f"{kind}_raw"] is not None]
        own = [r[kind] for r in rows if r[kind] is not None]
        out[kind] = scale.norm(kind, _mean(raw))
        out[f"{kind}_raw"] = _mean(raw)
        out[f"{kind}_count"] = len(raw)
        out[f"{kind}_image_mean"] = _mean(own)
        out[f"{kind}_min"] = min(own) if own else None
        out[f"{kind}_max"] = max(own) if own else None
        out[f"{kind}_stdev"] = statistics.pstdev(own) if len(own) > 1 else None
    out["overall"] = combine(out["face"], out["body"], face_weight, body_weight)
    own = [r["overall"] for r in rows if r["overall"] is not None]
    out["overall_image_mean"], out["overall_count"] = _mean(own), len(own)
    out["overall_min"], out["overall_max"] = (min(own), max(own)) if own else (None, None)
    out["overall_stdev"] = statistics.pstdev(own) if len(own) > 1 else None
    return out


def pct(v) -> str:
    return "n/a" if v is None or (isinstance(v, float) and math.isnan(v)) else f"{v:.1f}%"


def summary_text(summary: dict, refs: RefSet) -> str:
    n = summary["images"]
    face = f"Face {pct(summary['face'])}" + (
        f" ({summary['face_count']} of {n} with a face)" if refs.face else
        " (no face in any face reference)" if refs.face_info else " (no face references)")
    body = f"Body {pct(summary['body'])}" + (
        f" ({summary['body_count']} of {n} with a usable body)" if refs.body else
        " (no usable body in any body reference)" if refs.body_info else " (no body references)")
    lines = [f"Likeness of {n} image{'s' if n != 1 else ''}", f"Overall {pct(summary['overall'])}", face, body]
    if n > 1 and summary["overall_count"]:
        lines.append(f"Images' overall scores range from {pct(summary['overall_min'])} to "
                     f"{pct(summary['overall_max'])}")
    return "\n".join(lines)


def details_text(rows) -> str:
    out = []
    for r in rows:
        face = pct(r["face"]) if r["face"] is not None else ("no face" if r["faces"] == 0 else "n/a")
        body = pct(r["body"]) if r["body"] is not None else "n/a"
        out.append(f"{r['index']:>3}. {Path(r['image']).name}: overall {pct(r['overall'])} | face {face} | "
                   f"body {body}")
    return "\n".join(out)


# ---------------------------------------------------------------- spreadsheet

IMAGE_COLUMNS = [  # (key, header)
    ("index", "#"), ("image", "Image"), ("overall", "Overall"), ("face", "Face"), ("body", "Body"),
    ("faces", "Faces found"), ("best_face_ref", "Closest face reference"), ("face_cosine", "Face cosine"),
    ("face_raw", "Face raw %"), ("body_raw", "Body raw %"), ("body_compared", "Body references compared"),
    ("overall_raw", "Overall raw %"), ("path", "File")]
PERCENT_KEYS = {"overall", "face", "body", "face_raw", "body_raw", "overall_raw"}


def _summary_rows(summary, scale: Scale, refs: RefSet):
    yield ["", "Overall", "Face", "Body"]
    yield ["Average (as the trainer averages a probe)", summary["overall"], summary["face"], summary["body"]]
    yield ["Mean of the images' own scores", summary["overall_image_mean"], summary["face_image_mean"],
           summary["body_image_mean"]]
    yield ["Lowest image", summary["overall_min"], summary["face_min"], summary["body_min"]]
    yield ["Highest image", summary["overall_max"], summary["face_max"], summary["body_max"]]
    yield ["Standard deviation", summary["overall_stdev"], summary["face_stdev"], summary["body_stdev"]]
    yield ["Images scored", summary["overall_count"], summary["face_count"], summary["body_count"]]
    yield ["Raw average %", None, summary["face_raw"], summary["body_raw"]]
    yield []
    yield ["Images", summary["images"]]
    yield ["Face / body weight", f"{summary['face_weight']:g} / {summary['body_weight']:g}"]
    yield ["Normalized", "yes" if scale.normalize else "no"]
    for kind in ("face", "body"):
        lo, hi = getattr(scale, kind)
        yield [f"{kind.title()} scale (raw floor, ceiling)", lo, hi,
               f"floor: {scale.sources[kind]['floor']}, ceiling: {scale.sources[kind]['ceiling']}"]
    yield ["Face references (with a face, given)", refs.face_count, len(refs.face_info)]
    yield [f"Body references ({refs.body_kind}; usable, given)", refs.body_count, len(refs.body_info)]


def _reference_rows(refs: RefSet):
    yield ["Kind", "Reference", "Usable", "Details"]
    for info in refs.face_info:
        yield ["face", refs.label(info.get("image")), "yes" if info.get("faces") else "no face",
               "cut off by the frame" if info.get("cut_off") else ""]
    for info in refs.body_info:
        if refs.body_kind == "appearance":
            usable = ("yes" if info.get("used") else "no person" if not info.get("person")
                      else "body window not in frame")
            detail = f"window coverage {info['cover']:.2f}" if info.get("cover") is not None else ""
        else:
            usable = "yes" if info.get("ratios") else "no pose"
            detail = f"{info.get('ratios', 0)} ratios"
        yield ["body", refs.label(info.get("image")), usable, detail]


def _cell(v):
    if isinstance(v, float):
        return None if math.isnan(v) else round(v, 4)
    return v


def export(rows, summary, scale: Scale, refs: RefSet, base, fmt: str = "xlsx") -> list[Path]:
    """Write the scores: an .xlsx with Summary, Images and References sheets, or three CSV files (fmt "csv", or when
    openpyxl is missing). base: the path without an extension. Returns the files written."""
    base = Path(base)
    base.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "xlsx":
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            log.warning("openpyxl is not installed: writing CSV files instead of an .xlsx")
            fmt = "csv"
    if fmt == "csv":
        files = [base.with_name(base.name + "_images.csv"), base.with_name(base.name + "_summary.csv"),
                 base.with_name(base.name + "_references.csv")]
        with open(files[0], "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow([h for _, h in IMAGE_COLUMNS])
            for r in rows:
                w.writerow([_cell(r.get(k)) for k, _ in IMAGE_COLUMNS])
        for path, gen in ((files[1], _summary_rows(summary, scale, refs)), (files[2], _reference_rows(refs))):
            with open(path, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows([[_cell(v) for v in row] for row in gen])
        return files
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in _summary_rows(summary, scale, refs):
        ws.append([_cell(v) for v in row])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in range(2, 9):
        for c in range(2, 5):
            if isinstance(ws.cell(r, c).value, float):
                ws.cell(r, c).number_format = "0.0"
    ws.column_dimensions["A"].width = 46
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    wi = wb.create_sheet("Images")
    wi.append([h for _, h in IMAGE_COLUMNS])
    for r in rows:
        wi.append([_cell(r.get(k)) for k, _ in IMAGE_COLUMNS])
    for c in wi[1]:
        c.font = Font(bold=True)
    wi.freeze_panes = "C2"
    wi.auto_filter.ref = wi.dimensions
    for j, (key, header) in enumerate(IMAGE_COLUMNS, start=1):
        letter = get_column_letter(j)
        num = "0.0" if key in PERCENT_KEYS else "0.000" if key == "face_cosine" else None
        if num:
            for cell in wi[letter][1:]:
                cell.number_format = num
        longest = max([len(header)] + [len(str(_cell(r.get(key)) or "")) for r in rows])
        wi.column_dimensions[letter].width = min(60, max(8, longest + 2))
    wr = wb.create_sheet("References")
    for row in _reference_rows(refs):
        wr.append(row)
    for c in wr[1]:
        c.font = Font(bold=True)
    wr.column_dimensions["B"].width = 60
    wr.column_dimensions["C"].width = 24
    wr.column_dimensions["D"].width = 26
    path = base.with_name(base.name + ".xlsx")
    wb.save(path)
    return [path]
